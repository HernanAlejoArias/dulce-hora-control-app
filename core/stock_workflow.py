import os
import re
from datetime import date, datetime, timedelta
from typing import Any, Dict, List

import pandas as pd
from openpyxl import load_workbook

from .product_config import MovimientoStock
from .product_logic import (
    _generar_id_movimiento,
    descontar_stock_fefo,
    ingresar_entrega,
    load_json,
    save_json,
)

STATE_FILE = "proceso_stock_state.json"
DELIVERY_FILENAME_RE = re.compile(
    r"^Planilla de pedido Banfield\s*-\s*(\d{1,2})_(\d{1,2})_(\d{4})(?:\.[^.]+)?$",
    re.IGNORECASE,
)


def _archivos_dir() -> str:
    return os.environ.get(
        "DULCE_HORA_ARCHIVOS_DIR",
        os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "Archivos")),
    )


def _stock_file() -> str:
    return os.environ.get(
        "DULCE_HORA_STOCK_FILE",
        os.path.join(_archivos_dir(), "Stock", "Planilla de stock.xlsx"),
    )


def _ubicacion_file() -> str:
    return os.environ.get(
        "DULCE_HORA_UBICACION_STOCK_FILE",
        os.path.join(_archivos_dir(), "Stock", "Ubicacion Stock.xlsx"),
    )


def _ventas_file() -> str:
    return os.environ.get(
        "DULCE_HORA_VENTAS_FILE",
        os.path.join(_archivos_dir(), "Ventas", "Estadisticas de Banfield - 2026-06-01.xlsx"),
    )


def _empty_state() -> Dict[str, Any]:
    return {
        "conteo_stock": {"processed_dates": {}, "last_processed_date": None},
        "entradas": {"processed_files": {}, "last_processed_date": None},
        "ventas": {"processed_dates": {}, "last_processed_date": None},
        "desperdicio": {"processed_files": {}, "last_processed_date": None},
        "restauracion_stock": {"last_restored_at": None, "last_restored_count_date": None},
    }


def _ensure_state_shape(state: Dict[str, Any]) -> Dict[str, Any]:
    base = _empty_state()
    if not isinstance(state, dict):
        return base
    for section, defaults in base.items():
        if section not in state or not isinstance(state[section], dict):
            state[section] = defaults
            continue
        for key, val in defaults.items():
            state[section].setdefault(key, val)
    return state


def _infer_processed_state(state: Dict[str, Any]) -> Dict[str, Any]:
    movimientos = load_json("movimientos_stock.json")
    _hydrate_restore_info_from_forced_count(state)
    restore_at = _parse_datetime_value(state["restauracion_stock"].get("last_restored_at"))
    restore_count_date = _normalize_date_key(state["restauracion_stock"].get("last_restored_count_date"))
    for mov in movimientos:
        tipo = mov.get("tipo")
        archivo = mov.get("archivo_origen")
        obs = mov.get("observacion") or ""
        mov_dt = _parse_datetime_value(mov.get("fecha"))

        if tipo == "RECEPCION" and archivo:
            entrega_fecha = _delivery_date_from_filename_safe(archivo)
            if _should_ignore_inferred_process(entrega_fecha, mov_dt, restore_count_date, restore_at):
                continue
            state["entradas"]["processed_files"].setdefault(archivo, {"inferido": True})

        if tipo == "VENTA":
            match = re.search(r"Venta del (\d{4}-\d{2}-\d{2})", obs)
            if match:
                fecha = match.group(1)
                if _should_ignore_inferred_process(fecha, mov_dt, restore_count_date, restore_at):
                    continue
                state["ventas"]["processed_dates"].setdefault(fecha, {"inferido": True})

        if tipo == "DESPERDICIO":
            match = re.search(r"Desperdicio txt (\d{8})", obs)
            if match:
                archivo_txt = f"{match.group(1)}.txt"
                fecha = _normalize_date_key(match.group(1))
                if _should_ignore_inferred_process(fecha, mov_dt, restore_count_date, restore_at):
                    continue
                state["desperdicio"]["processed_files"].setdefault(archivo_txt, {"inferido": True})

        if tipo in {"AJUSTE_POSITIVO", "AJUSTE_NEGATIVO"}:
            match = re.search(r"Conteo (\d{4}-\d{2}-\d{2})", obs)
            if match:
                fecha = match.group(1)
                state["conteo_stock"]["processed_dates"].setdefault(fecha, {"inferido": True})

    for section, collection_key in [
        ("conteo_stock", "processed_dates"),
        ("ventas", "processed_dates"),
        ("desperdicio", "processed_files"),
    ]:
        keys = sorted(state[section][collection_key].keys())
        state[section]["last_processed_date"] = keys[-1] if keys else state[section].get("last_processed_date")

    entrada_fechas = [
        info.get("fecha")
        for info in state["entradas"]["processed_files"].values()
        if isinstance(info, dict) and info.get("fecha")
    ]
    if entrada_fechas:
        state["entradas"]["last_processed_date"] = sorted(entrada_fechas)[-1]

    return state


def load_process_state() -> Dict[str, Any]:
    state = _ensure_state_shape(load_json(STATE_FILE, dict))
    return _infer_processed_state(state)


def save_process_state(state: Dict[str, Any]) -> None:
    save_json(STATE_FILE, _ensure_state_shape(state))


def _parse_datetime_value(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        return value
    if not value:
        return None
    try:
        return datetime.fromisoformat(str(value))
    except ValueError:
        return None


def _normalize_date_key(value: Any) -> str | None:
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if not value:
        return None
    text = str(value).strip()
    if re.fullmatch(r"\d{8}", text):
        return f"{text[:4]}-{text[4:6]}-{text[6:8]}"
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", text):
        return text
    return None


def _delivery_date_from_filename_safe(filename: str) -> str | None:
    match = DELIVERY_FILENAME_RE.match(os.path.basename(filename))
    if not match:
        return None
    day, month, year = match.groups()
    return f"{int(year):04d}-{int(month):02d}-{int(day):02d}"


def _should_ignore_inferred_process(
    process_date: str | None,
    movement_dt: datetime | None,
    restore_count_date: str | None,
    restore_at: datetime | None,
) -> bool:
    if not process_date or not movement_dt or not restore_count_date or not restore_at:
        return False
    return process_date >= restore_count_date and movement_dt <= restore_at


def _hydrate_restore_info_from_forced_count(state: Dict[str, Any]) -> None:
    if state["restauracion_stock"].get("last_restored_at"):
        return

    forced = []
    for fecha, info in state["conteo_stock"]["processed_dates"].items():
        if not isinstance(info, dict) or not info.get("forzado") or not info.get("processed_at"):
            continue
        processed_at = _parse_datetime_value(info.get("processed_at"))
        fecha_key = _normalize_date_key(fecha)
        if processed_at and fecha_key:
            forced.append((processed_at, fecha_key))

    if not forced:
        return

    restored_at, count_date = sorted(forced, key=lambda item: item[0])[-1]
    state["restauracion_stock"] = {
        "last_restored_at": restored_at.isoformat(),
        "last_restored_count_date": count_date,
    }
    _clear_reprocess_flags_after_restore(state, count_date)


def _recompute_last_processed_dates(state: Dict[str, Any]) -> None:
    entradas_fechas = [
        info.get("fecha") or _delivery_date_from_filename_safe(archivo)
        for archivo, info in state["entradas"]["processed_files"].items()
        if isinstance(info, dict)
    ]
    entradas_fechas = [fecha for fecha in entradas_fechas if fecha]
    state["entradas"]["last_processed_date"] = sorted(entradas_fechas)[-1] if entradas_fechas else None

    ventas_fechas = sorted(state["ventas"]["processed_dates"].keys())
    state["ventas"]["last_processed_date"] = ventas_fechas[-1] if ventas_fechas else None

    desperdicio_fechas = [
        _normalize_date_key(info.get("fecha")) or _normalize_date_key(archivo.replace(".txt", ""))
        for archivo, info in state["desperdicio"]["processed_files"].items()
        if isinstance(info, dict)
    ]
    desperdicio_fechas = [fecha for fecha in desperdicio_fechas if fecha]
    state["desperdicio"]["last_processed_date"] = sorted(desperdicio_fechas)[-1] if desperdicio_fechas else None


def _clear_reprocess_flags_after_restore(state: Dict[str, Any], count_date: str) -> Dict[str, List[str]]:
    removed = {"entradas": [], "ventas": [], "desperdicio": []}

    for archivo, info in list(state["entradas"]["processed_files"].items()):
        fecha = info.get("fecha") if isinstance(info, dict) else None
        fecha = _normalize_date_key(fecha) or _delivery_date_from_filename_safe(archivo)
        if fecha and fecha >= count_date:
            del state["entradas"]["processed_files"][archivo]
            removed["entradas"].append(archivo)

    for fecha in list(state["ventas"]["processed_dates"].keys()):
        fecha_key = _normalize_date_key(fecha)
        if fecha_key and fecha_key >= count_date:
            del state["ventas"]["processed_dates"][fecha]
            removed["ventas"].append(fecha)

    for archivo, info in list(state["desperdicio"]["processed_files"].items()):
        fecha = info.get("fecha") if isinstance(info, dict) else None
        fecha = _normalize_date_key(fecha) or _normalize_date_key(archivo.replace(".txt", ""))
        if fecha and fecha >= count_date:
            del state["desperdicio"]["processed_files"][archivo]
            removed["desperdicio"].append(archivo)

    _recompute_last_processed_dates(state)
    return removed


def _safe_float(value: Any) -> float | None:
    if value is None or value == "":
        return None
    try:
        parsed = float(value)
    except (TypeError, ValueError):
        return None
    if pd.isna(parsed):
        return None
    return parsed


def _format_qty(value: float) -> float | int:
    rounded = round(float(value), 3)
    return int(rounded) if rounded.is_integer() else rounded


def _parse_count_date(value: str | None) -> datetime:
    raw = str(value or "").strip()
    for fmt in ("%Y%m%d", "%Y-%m-%d"):
        try:
            return datetime.strptime(raw, fmt)
        except ValueError:
            pass
    raise ValueError("La fecha del conteo debe tener formato aaaammdd.")


def _stock_actual() -> Dict[str, float]:
    lotes = load_json("stock_lotes.json")
    stock: Dict[str, float] = {}
    for lote in lotes:
        if lote.get("estado") == "ACTIVO" and _safe_float(lote.get("cantidad_actual")) and lote.get("cantidad_actual") > 0:
            codigo = str(lote.get("codigo_producto"))
            stock[codigo] = stock.get(codigo, 0.0) + float(lote.get("cantidad_actual", 0))
    return stock


def load_ubicaciones() -> Dict[str, Dict[str, str]]:
    path = _ubicacion_file()
    ubicaciones: Dict[str, Dict[str, str]] = {}
    if not os.path.exists(path):
        return ubicaciones

    df = pd.read_excel(path, header=0)
    for col in df.columns[1:]:
        codigo = str(col).strip()
        if not codigo.isdigit():
            continue
        articulo = str(df.iloc[0][col]).strip() if len(df) and pd.notna(df.iloc[0][col]) else ""
        ubicacion = "Sin ubicacion"
        for idx in range(1, len(df)):
            val = df.iloc[idx][col]
            if pd.notna(val):
                ubicacion = str(val).strip()
                break
        ubicaciones[codigo] = {"ubicacion": ubicacion, "articulo_planilla": articulo}
    return ubicaciones


def load_latest_stock_count() -> Dict[str, Any] | None:
    path = _stock_file()
    if not os.path.exists(path):
        return None

    df = pd.read_excel(path, header=0)
    fechas: List[tuple[int, datetime]] = []
    for idx, value in df.iloc[:, 0].items():
        parsed = pd.to_datetime(value, errors="coerce")
        if pd.notna(parsed):
            fechas.append((idx, parsed.to_pydatetime()))

    if not fechas:
        return None

    fechas.sort(key=lambda item: item[1], reverse=True)
    row_idx, fecha = fechas[0]
    row = df.iloc[row_idx]
    valores: Dict[str, float] = {}
    orden: List[str] = []
    for col in df.columns[1:]:
        codigo = str(col).strip()
        if not codigo.isdigit():
            continue
        cantidad = _safe_float(row[col])
        if cantidad is None:
            continue
        valores[codigo] = cantidad
        orden.append(codigo)

    fecha_str = fecha.strftime("%Y-%m-%d")
    return {
        "fecha": fecha_str,
        "archivo": os.path.basename(path),
        "path": path,
        "valores": valores,
        "orden": orden,
        "es_futura": fecha.date() > date.today(),
    }


def guardar_conteo_stock(fecha: str, valores: Dict[str, Any]) -> Dict[str, Any]:
    path = _stock_file()
    if not os.path.exists(path):
        return {"status": "error", "message": "No se encontro la planilla de stock."}

    fecha_dt = _parse_count_date(fecha)
    wb = load_workbook(path)
    ws = wb.active
    codigo_col: Dict[str, int] = {}
    for col_idx in range(2, ws.max_column + 1):
        codigo = str(ws.cell(row=1, column=col_idx).value or "").strip()
        if codigo.isdigit():
            codigo_col[codigo] = col_idx

    target_row = None
    first_empty_row = None
    for row_idx in range(3, ws.max_row + 1):
        cell = ws.cell(row=row_idx, column=1)
        if cell.value is None and first_empty_row is None:
            first_empty_row = row_idx
            continue
        parsed = pd.to_datetime(cell.value, errors="coerce")
        if pd.notna(parsed) and parsed.to_pydatetime().date() == fecha_dt.date():
            target_row = row_idx
            break

    if target_row is None:
        target_row = first_empty_row or (ws.max_row + 1)

    date_cell = ws.cell(row=target_row, column=1)
    date_cell.value = fecha_dt
    date_cell.number_format = "yyyymmdd"

    guardados = 0
    omitidos = []
    for codigo, value in valores.items():
        codigo_str = str(codigo).strip()
        col_idx = codigo_col.get(codigo_str)
        if not col_idx:
            omitidos.append(codigo_str)
            continue
        cantidad = _safe_float(value)
        ws.cell(row=target_row, column=col_idx).value = cantidad
        guardados += 1

    wb.save(path)
    return {
        "status": "ok",
        "fecha": fecha_dt.strftime("%Y-%m-%d"),
        "fecha_formato": fecha_dt.strftime("%Y%m%d"),
        "archivo": os.path.basename(path),
        "fila": target_row,
        "guardados": guardados,
        "omitidos": omitidos,
    }


def _stock_cutoff_date(state: Dict[str, Any] | None = None) -> str | None:
    state = state or load_process_state()
    fechas = set(state["conteo_stock"]["processed_dates"].keys())
    latest_count = load_latest_stock_count()
    if latest_count and latest_count.get("fecha"):
        fechas.add(latest_count["fecha"])
    return sorted(fechas)[-1] if fechas else None


def get_conteo_planilla() -> Dict[str, Any]:
    productos = load_json("productos.json")
    prod_dict = {str(p.get("codigo_producto")): p for p in productos if p.get("activo", True)}
    ubicaciones = load_ubicaciones()
    stock = _stock_actual()
    conteo = load_latest_stock_count()
    state = load_process_state()
    processed_dates = state["conteo_stock"]["processed_dates"]

    grupos: Dict[str, List[Dict[str, Any]]] = {}
    orden = conteo["orden"] if conteo else list(prod_dict.keys())
    codigos = orden + [cod for cod in prod_dict if cod not in orden]

    for codigo in codigos:
        prod = prod_dict.get(codigo)
        if not prod:
            continue
        ubicacion = ubicaciones.get(codigo, {}).get("ubicacion", "Sin ubicacion")
        contado = conteo["valores"].get(codigo) if conteo else None
        actual = stock.get(codigo, 0.0)
        item = {
            "codigo": codigo,
            "descripcion": prod.get("descripcion", ubicaciones.get(codigo, {}).get("articulo_planilla", "")),
            "categoria": prod.get("categoria", ""),
            "ubicacion": ubicacion,
            "stock_actual": _format_qty(actual),
            "stock_planilla": _format_qty(contado) if contado is not None else None,
            "diferencia": _format_qty(contado - actual) if contado is not None else None,
        }
        grupos.setdefault(ubicacion, []).append(item)

    return {
        "conteo": {
            "fecha": conteo["fecha"] if conteo else None,
            "archivo": conteo["archivo"] if conteo else None,
            "cantidad_productos": len(conteo["valores"]) if conteo else 0,
            "procesado": bool(conteo and conteo["fecha"] in processed_dates),
            "es_futura": bool(conteo and conteo["es_futura"]),
        },
        "grupos": [{"ubicacion": key, "productos": val} for key, val in sorted(grupos.items())],
    }


def _append_movimiento(movimientos: List[Dict[str, Any]], **kwargs: Any) -> None:
    mov = MovimientoStock(
        id_movimiento=_generar_id_movimiento(),
        fecha=datetime.now(),
        **kwargs,
    )
    movimientos.append(mov.model_dump(mode="json"))


def _unique_lote_id(lotes: List[Dict[str, Any]], base_id: str) -> str:
    existing = {l.get("id_lote") for l in lotes}
    if base_id not in existing:
        return base_id
    suffix = 2
    while f"{base_id}-{suffix}" in existing:
        suffix += 1
    return f"{base_id}-{suffix}"


def procesar_conteo_stock(forzar: bool = False) -> Dict[str, Any]:
    conteo = load_latest_stock_count()
    if not conteo:
        return {"status": "error", "message": "No se encontro una fila de conteo en la planilla de stock."}

    state = load_process_state()
    fecha = conteo["fecha"]
    if not forzar and fecha in state["conteo_stock"]["processed_dates"]:
        return {"status": "skipped", "message": f"El conteo {fecha} ya fue procesado.", "fecha": fecha}

    productos = load_json("productos.json")
    prod_dict = {str(p.get("codigo_producto")): p for p in productos}
    stock = _stock_actual()
    total_pos = 0.0
    total_neg = 0.0
    total_conteos = 0
    archivo = conteo["archivo"]
    accion = "Restauracion" if forzar else "Ajuste"

    for codigo, contado in conteo["valores"].items():
        actual = stock.get(codigo, 0.0)
        diff = round(contado - actual, 3)
        if diff < 0:
            descontar_stock_fefo(
                codigo,
                abs(diff),
                "AJUSTE_NEGATIVO",
                f"{accion} x Conteo {fecha}",
                archivo_origen=archivo,
            )
            total_neg += abs(diff)

    lotes = load_json("stock_lotes.json")
    movimientos = load_json("movimientos_stock.json")
    fecha_dt = datetime.fromisoformat(fecha)

    for codigo, contado in conteo["valores"].items():
        actual = stock.get(codigo, 0.0)
        diff = round(contado - actual, 3)
        prod = prod_dict.get(codigo, {})

        _append_movimiento(
            movimientos,
            codigo_producto=codigo,
            id_lote=None,
            tipo="CONTEO_STOCK",
            cantidad=contado,
            observacion=f"{accion} fisico {fecha}",
            archivo_origen=archivo,
        )
        total_conteos += 1

        if diff <= 0:
            continue

        vencimiento = None
        vida_util = prod.get("vida_util_dias")
        if vida_util and vida_util > 0:
            vencimiento = (fecha_dt + timedelta(days=int(vida_util))).strftime("%Y-%m-%d")

        base_lote_id = f"{codigo}-CONTEO-{fecha_dt.strftime('%Y%m%d')}"
        lote_id = _unique_lote_id(lotes, base_lote_id)
        nuevo_lote = {
            "id_lote": lote_id,
            "codigo_producto": codigo,
            "lote": fecha_dt.strftime("%Y%m%d"),
            "descripcion_producto": prod.get("descripcion", "Desconocido"),
            "categoria": prod.get("categoria", "Sin Categoria"),
            "cantidad_inicial": diff,
            "cantidad_actual": diff,
            "fecha_ingreso": fecha,
            "fecha_vencimiento": vencimiento,
            "estado": "ACTIVO",
            "origen": f"{accion} Positivo Conteo {fecha}",
            "archivo_origen": archivo,
        }
        lotes.append(nuevo_lote)
        _append_movimiento(
            movimientos,
            codigo_producto=codigo,
            id_lote=lote_id,
            tipo="AJUSTE_POSITIVO",
            cantidad=diff,
            observacion=f"{accion} Conteo {fecha}",
            archivo_origen=archivo,
        )
        total_pos += diff

    save_json("stock_lotes.json", lotes)
    save_json("movimientos_stock.json", movimientos)

    processed_at = datetime.now().isoformat()
    reprocess_flags_cleared = {"entradas": [], "ventas": [], "desperdicio": []}
    if forzar:
        state["restauracion_stock"] = {
            "last_restored_at": processed_at,
            "last_restored_count_date": fecha,
        }
        reprocess_flags_cleared = _clear_reprocess_flags_after_restore(state, fecha)

    state["conteo_stock"]["processed_dates"][fecha] = {
        "processed_at": processed_at,
        "archivo": archivo,
        "ajustes_positivos": _format_qty(total_pos),
        "ajustes_negativos": _format_qty(total_neg),
        "conteos": total_conteos,
        "forzado": forzar,
    }
    state["conteo_stock"]["last_processed_date"] = fecha
    save_process_state(state)

    return {
        "status": "ok",
        "fecha": fecha,
        "conteos": total_conteos,
        "ajustes_positivos": _format_qty(total_pos),
        "ajustes_negativos": _format_qty(total_neg),
        "forzado": forzar,
        "flags_liberados": reprocess_flags_cleared,
    }


def _delivery_date_from_filename(path: str) -> datetime:
    filename = os.path.basename(path)
    match = DELIVERY_FILENAME_RE.match(filename)
    if not match:
        raise ValueError(
            "El archivo de pedido debe llamarse 'Planilla de pedido Banfield - d_m_aaaa.xlsx'."
        )
    day, month, year = match.groups()
    return datetime(int(year), int(month), int(day))


def _parse_delivery_file(path: str) -> Dict[str, Any]:
    df = pd.read_excel(path)
    items: List[Dict[str, Any]] = []
    for idx, row in df.iterrows():
        if idx < 2:
            continue
        cod = str(row.iloc[0]).strip()
        if not cod or cod == "nan" or not cod.isdigit() or cod in {"1011", "999"}:
            continue
        cantidad = _safe_float(row.iloc[4] if len(row) > 4 else None)
        if cantidad and cantidad > 0:
            items.append({"codigo": str(int(float(cod))), "cantidad": cantidad})
    return {
        "archivo": os.path.basename(path),
        "fecha": _delivery_date_from_filename(path).strftime("%Y-%m-%d"),
        "items": items,
    }


def list_delivery_files() -> List[Dict[str, Any]]:
    entregas_dir = os.path.join(_archivos_dir(), "Entregas")
    state = load_process_state()
    processed = state["entradas"]["processed_files"]
    cutoff = _stock_cutoff_date(state)
    if not os.path.isdir(entregas_dir):
        return []
    files = []
    for name in sorted(os.listdir(entregas_dir)):
        if not name.lower().endswith((".xlsx", ".xlsm", ".xls")):
            continue
        if not DELIVERY_FILENAME_RE.match(name):
            continue
        info = _parse_delivery_file(os.path.join(entregas_dir, name))
        info["procesado"] = name in processed
        info["bloqueado_por_conteo"] = bool(cutoff and info["fecha"] < cutoff and not info["procesado"])
        info["fecha_corte_stock"] = cutoff
        files.append(info)
    return files


def procesar_entradas(archivos_seleccionados: List[str] | None = None) -> Dict[str, Any]:
    state = load_process_state()
    processed = state["entradas"]["processed_files"]
    seleccion = set(archivos_seleccionados) if archivos_seleccionados is not None else None
    procesados = []
    omitidos = []
    bloqueados = []
    for info in list_delivery_files():
        archivo = info["archivo"]
        if seleccion is not None and archivo not in seleccion:
            continue
        if archivo in processed:
            omitidos.append(archivo)
            continue
        if info.get("bloqueado_por_conteo"):
            bloqueados.append({
                "archivo": archivo,
                "fecha": info["fecha"],
                "fecha_corte_stock": info.get("fecha_corte_stock"),
                "motivo": "La entrega es anterior al ultimo conteo de stock.",
            })
            continue
        if not info["items"]:
            omitidos.append(archivo)
            continue
        ingresar_entrega(info["items"], datetime.fromisoformat(info["fecha"]), archivo)
        processed[archivo] = {
            "processed_at": datetime.now().isoformat(),
            "fecha": info["fecha"],
            "items": len(info["items"]),
            "cantidad_total": _format_qty(sum(i["cantidad"] for i in info["items"])),
        }
        state["entradas"]["last_processed_date"] = info["fecha"]
        procesados.append(processed[archivo] | {"archivo": archivo})
    save_process_state(state)
    return {"status": "ok", "procesados": procesados, "omitidos": omitidos, "bloqueados": bloqueados}


def _sales_rows() -> List[Dict[str, Any]]:
    path = _ventas_file()
    if not os.path.exists(path):
        return []
    df = pd.read_excel(path, sheet_name="Estadisticas de Productos", header=None)
    codigos: Dict[int, str] = {}
    for col_idx, cod in df.iloc[3].to_dict().items():
        parsed = _safe_float(cod)
        if parsed is not None:
            codigos[col_idx] = str(int(parsed))

    rows = []
    for idx, row in df.iterrows():
        if idx < 8:
            continue
        fecha = pd.to_datetime(row.iloc[0], errors="coerce")
        if pd.isna(fecha):
            continue
        items = []
        for col_idx, codigo in codigos.items():
            cantidad = _safe_float(row.iloc[col_idx])
            if cantidad and cantidad > 0:
                items.append({"codigo": codigo, "cantidad": cantidad})
        if items:
            rows.append({"fecha": fecha.strftime("%Y-%m-%d"), "items": items})
    return rows


def _encontrar_codigo_por_texto(texto: str, prod_dict: Dict[str, Dict[str, Any]]) -> str | None:
    texto_norm = texto.lower()
    for codigo, prod in prod_dict.items():
        desc = str(prod.get("descripcion", "")).lower()
        if texto_norm in desc or desc in texto_norm:
            return codigo
        palabras = texto_norm.split()
        if len(palabras) > 1 and palabras[0] in desc and palabras[1] in desc:
            return codigo
    manual = {
        "medialunas": "100",
        "surtidas": "102",
        "pan frances": "200",
        "pizza muzzarella": "500",
        "pote chocotorta": "301",
        "bizcocho grasa": "206",
    }
    for key, codigo in manual.items():
        if key in texto_norm:
            return codigo
    return None


def _waste_files() -> List[Dict[str, Any]]:
    desperdicio_dir = os.path.join(_archivos_dir(), "Desperdicio")
    if not os.path.isdir(desperdicio_dir):
        return []
    productos = load_json("productos.json")
    prod_dict = {str(p.get("codigo_producto")): p for p in productos}
    files = []
    for name in sorted(os.listdir(desperdicio_dir)):
        if not name.lower().endswith(".txt"):
            continue
        path = os.path.join(desperdicio_dir, name)
        items = []
        with open(path, "r", encoding="utf-8", errors="ignore") as handle:
            for line in handle.readlines():
                text = line.strip()
                if not text or text.startswith("ID") or "Eliminar desperdicio" in text or text.startswith("$"):
                    continue
                match = re.search(r"(.+?)\s+(\d+(?:[\.,]\d+)?)\s+\$", text)
                if not match:
                    continue
                codigo = _encontrar_codigo_por_texto(match.group(1).strip(), prod_dict)
                cantidad = _safe_float(match.group(2).replace(",", "."))
                if codigo and cantidad and cantidad > 0:
                    items.append({"codigo": codigo, "cantidad": cantidad, "texto": match.group(1).strip()})
        files.append({"archivo": name, "fecha": name.replace(".txt", ""), "items": items})
    return files


def procesar_ventas_desperdicio() -> Dict[str, Any]:
    state = load_process_state()
    ventas_processed = state["ventas"]["processed_dates"]
    waste_processed = state["desperdicio"]["processed_files"]
    ventas_ok = []
    ventas_omitidas = []
    desperdicio_ok = []
    desperdicio_omitido = []
    ventas_archivo = os.path.basename(_ventas_file())

    for row in _sales_rows():
        fecha = row["fecha"]
        if fecha in ventas_processed:
            ventas_omitidas.append(fecha)
            continue
        total = 0.0
        for item in row["items"]:
            descontar_stock_fefo(
                item["codigo"],
                item["cantidad"],
                "VENTA",
                f"Venta del {fecha}",
                archivo_origen=ventas_archivo,
            )
            total += item["cantidad"]
        ventas_processed[fecha] = {
            "processed_at": datetime.now().isoformat(),
            "archivo": ventas_archivo,
            "items": len(row["items"]),
            "cantidad_total": _format_qty(total),
        }
        state["ventas"]["last_processed_date"] = fecha
        ventas_ok.append({"fecha": fecha, **ventas_processed[fecha]})

    for info in _waste_files():
        archivo = info["archivo"]
        if archivo in waste_processed:
            desperdicio_omitido.append(archivo)
            continue
        total = 0.0
        fecha = info["fecha"]
        for item in info["items"]:
            descontar_stock_fefo(
                item["codigo"],
                item["cantidad"],
                "DESPERDICIO",
                f"Desperdicio txt {fecha}",
                archivo_origen=archivo,
            )
            total += item["cantidad"]
        waste_processed[archivo] = {
            "processed_at": datetime.now().isoformat(),
            "fecha": fecha,
            "items": len(info["items"]),
            "cantidad_total": _format_qty(total),
        }
        state["desperdicio"]["last_processed_date"] = fecha
        desperdicio_ok.append({"archivo": archivo, **waste_processed[archivo]})

    save_process_state(state)
    return {
        "status": "ok",
        "ventas_procesadas": ventas_ok,
        "ventas_omitidas": ventas_omitidas,
        "desperdicio_procesado": desperdicio_ok,
        "desperdicio_omitido": desperdicio_omitido,
    }


def get_stock_workflow_status() -> Dict[str, Any]:
    stock = _stock_actual()
    conteo = get_conteo_planilla()["conteo"]
    deliveries = list_delivery_files()
    state = load_process_state()
    sales_rows = _sales_rows()
    waste_files = _waste_files()
    ventas_processed = state["ventas"]["processed_dates"]
    waste_processed = state["desperdicio"]["processed_files"]
    entradas_pendientes = [d for d in deliveries if not d["procesado"] and not d.get("bloqueado_por_conteo")]
    entradas_bloqueadas = [d for d in deliveries if d.get("bloqueado_por_conteo")]
    return {
        "stock_actual": {
            "productos_con_stock": sum(1 for qty in stock.values() if qty > 0),
            "unidades_totales": _format_qty(sum(stock.values())),
        },
        "conteo": conteo,
        "entradas": {
            "pendientes": entradas_pendientes,
            "bloqueadas": entradas_bloqueadas,
            "procesadas": [d for d in deliveries if d["procesado"]],
            "ultima_fecha_procesada": state["entradas"]["last_processed_date"],
            "fecha_corte_stock": _stock_cutoff_date(state),
        },
        "ventas_desperdicio": {
            "ventas_pendientes": [r for r in sales_rows if r["fecha"] not in ventas_processed],
            "ventas_procesadas": sorted(ventas_processed.keys()),
            "desperdicio_pendiente": [w for w in waste_files if w["archivo"] not in waste_processed],
            "desperdicio_procesado": sorted(waste_processed.keys()),
            "ultima_venta_procesada": state["ventas"]["last_processed_date"],
            "ultimo_desperdicio_procesado": state["desperdicio"]["last_processed_date"],
        },
    }
