from __future__ import annotations

import os
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal
from typing import Any, Dict, List

from openpyxl import load_workbook

CASHFLOW_SHEET = "Hoja 1"
CASHFLOW_HEADERS = [
    "Mes Imputacion",
    "Fecha vencimiento",
    "Categoria",
    "Ref1",
    "Destinatario",
    "Ref2",
    "Costo",
    "Fecha de pago",
    "M. Pago",
    "Origen",
    "Estado",
    "Facturado",
    "Notas",
    "Auxiliar",
]


def _archivos_dir() -> str:
    return os.environ.get(
        "DULCE_HORA_ARCHIVOS_DIR",
        os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "Archivos")),
    )


def cashflow_file() -> str:
    return os.environ.get(
        "DULCE_HORA_CASHFLOW_FILE",
        os.path.join(_archivos_dir(), "DH - CashFlow.xlsx"),
    )


def _load_workbook():
    path = cashflow_file()
    if not os.path.exists(path):
        raise FileNotFoundError(f"No se encontro el archivo de cashflow: {path}")
    workbook = load_workbook(path)
    if CASHFLOW_SHEET not in workbook.sheetnames:
        raise ValueError(f"El archivo no tiene la hoja esperada '{CASHFLOW_SHEET}'.")
    sheet = workbook[CASHFLOW_SHEET]
    headers = [sheet.cell(1, col).value for col in range(1, len(CASHFLOW_HEADERS) + 1)]
    if headers != CASHFLOW_HEADERS:
        raise ValueError("El archivo de cashflow no tiene las columnas esperadas.")
    return workbook


def _parse_yyyymmdd(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if value in (None, ""):
        return None
    try:
        return datetime.strptime(str(int(value)), "%Y%m%d").date()
    except (TypeError, ValueError):
        return None


def _parse_date_input(value: Any, field_name: str) -> date:
    parsed = _parse_yyyymmdd(value)
    if parsed:
        return parsed
    text = str(value or "").strip()
    for fmt in ("%Y-%m-%d", "%Y%m%d", "%d/%m/%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"{field_name} debe tener formato aaaammdd.")


def _parse_month(value: Any, payment_date: date) -> int:
    if value in (None, ""):
        return int(payment_date.strftime("%Y%m"))
    text = str(value).strip()
    if len(text) == 6 and text.isdigit():
        return int(text)
    raise ValueError("Mes Imputacion debe tener formato aaaamm.")


def _parse_amount(value: Any) -> float:
    if value in (None, ""):
        raise ValueError("Costo es obligatorio.")
    text = str(value).replace("$", "").replace(" ", "")
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        text = text.replace(".", "").replace(",", ".")
    try:
        amount = float(text)
    except ValueError as exc:
        raise ValueError("Costo debe ser numerico.") from exc
    if amount <= 0:
        raise ValueError("Costo debe ser mayor a cero.")
    return amount


def _row_to_movement(row: tuple[Any, ...], row_number: int) -> Dict[str, Any]:
    payment_date = _parse_yyyymmdd(row[7])
    due_date = _parse_yyyymmdd(row[1])
    amount = row[6] or 0
    return {
        "fila": row_number,
        "mes_imputacion": row[0],
        "fecha_vencimiento": due_date.strftime("%Y%m%d") if due_date else None,
        "categoria": row[2] or "",
        "ref1": row[3] or "",
        "destinatario": row[4] or "",
        "ref2": row[5] or "",
        "costo": float(amount) if amount else 0,
        "fecha_pago": payment_date.strftime("%Y%m%d") if payment_date else None,
        "medio_pago": row[8] or "",
        "origen": row[9] or "",
        "estado": row[10] or "",
        "facturado": row[11] or "",
        "notas": row[12] or "",
    }


def list_cashflow_movements(limit: int = 80) -> Dict[str, Any]:
    workbook = _load_workbook()
    sheet = workbook[CASHFLOW_SHEET]
    movements = []
    monthly: dict[int, Decimal] = defaultdict(lambda: Decimal("0"))
    categories: dict[str, Decimal] = defaultdict(lambda: Decimal("0"))

    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not any(row[:13]):
            continue
        movement = _row_to_movement(row, row_number)
        movements.append(movement)
        amount = Decimal(str(movement["costo"] or 0))
        if movement["mes_imputacion"]:
            monthly[int(movement["mes_imputacion"])] += amount
        categories[movement["categoria"] or "Sin categoria"] += amount

    recent = sorted(
        movements,
        key=lambda item: (str(item.get("fecha_pago") or ""), item["fila"]),
        reverse=True,
    )[:limit]
    latest_month = max(monthly.keys()) if monthly else None

    return {
        "archivo": os.path.basename(cashflow_file()),
        "ruta": cashflow_file(),
        "total_movimientos": len(movements),
        "ultimo_mes": latest_month,
        "total_ultimo_mes": float(monthly[latest_month]) if latest_month else 0,
        "total_general": float(sum(monthly.values(), Decimal("0"))),
        "resumen_mensual": [
            {"mes": month, "total": float(total)}
            for month, total in sorted(monthly.items(), reverse=True)
        ],
        "categorias": [
            {"categoria": category, "total": float(total)}
            for category, total in sorted(categories.items(), key=lambda item: item[1], reverse=True)
        ][:20],
        "movimientos": recent,
    }


def append_cashflow_movement(payload: Dict[str, Any]) -> Dict[str, Any]:
    workbook = _load_workbook()
    sheet = workbook[CASHFLOW_SHEET]

    payment_date = _parse_date_input(payload.get("fecha_pago"), "Fecha de pago")
    due_date = _parse_date_input(payload.get("fecha_vencimiento") or payload.get("fecha_pago"), "Fecha vencimiento")
    imputation_month = _parse_month(payload.get("mes_imputacion"), payment_date)
    amount = _parse_amount(payload.get("costo"))
    row_number = sheet.max_row + 1

    row = [
        imputation_month,
        int(due_date.strftime("%Y%m%d")),
        payload.get("categoria") or "Sin categoria",
        payload.get("ref1") or None,
        payload.get("destinatario") or None,
        payload.get("ref2") or None,
        amount,
        int(payment_date.strftime("%Y%m%d")),
        payload.get("medio_pago") or "Transferencia",
        payload.get("origen") or "Dulce Hora",
        payload.get("estado") or "Pago",
        payload.get("facturado") or None,
        payload.get("notas") or "Registrado desde app Dulce Hora",
        f"=CONCATENATE(A{row_number},B{row_number},H{row_number})",
    ]
    sheet.append(row)
    workbook.save(cashflow_file())

    return {
        "status": "ok",
        "fila": row_number,
        "movimiento": _row_to_movement(tuple(row), row_number),
    }
